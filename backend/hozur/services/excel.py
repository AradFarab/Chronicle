"""Excel import (ورود اکسلی): for every screen where the admin can add something there is a
sample workbook to download and an import that reads the filled workbook back.

The sample has one sheet to fill (a header row, formats already applied to the columns: text
columns stay text so "1405/07/14" is never turned into a Gregorian date, drop-down lists for
choices and for things that already exist), a guide sheet that says what each column wants,
and a hidden sheet with the lists. The import is all or nothing: every row goes through the same
code as the form (same validation, audit and recalculation), and if any row is wrong nothing is
saved and each wrong row is reported with its Excel row number.
"""

from __future__ import annotations

import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from hozur.calendar import service as calendar_service
from hozur.db.models import (
	ApprovalProcess,
	AttendanceDevice,
	Department,
	Designation,
	Employee,
	LeaveType,
	Province,
	Shift,
	WorkLocation,
	WorkPattern,
)
from hozur.engine.rules import PARAMETERS, baseline_rule_set
from hozur.i18n import _, translate
from hozur.jalali.dates import (
	WEEKDAY_NAMES,
	InvalidJalaliDate,
	normalize_jdate,
	normalize_time,
	to_ascii_digits,
)
from hozur.services import devices, employees, resources
from hozur.services.errors import HozurError, NotFound, ValidationError

MAX_ROWS = 5000
FORMATTED_ROWS = 1000  # rows of the sample sheet that carry formats and drop-downs

TEXT, NUMBER, BOOL, JDATE, TIME, CHOICE, REF = "text", "number", "bool", "jdate", "time", "choice", "ref"
YES_WORDS = {"1", "true", "yes", "y", "بله", "آری", "ماری"}
NO_WORDS = {"0", "false", "no", "n", "خیر", "نه"}


@dataclass(frozen=True)
class Col:
	key: str
	label: str  # English source text, translated into the header
	kind: str = TEXT
	required: bool = False
	choices: tuple[tuple[str, str], ...] = ()  # (code, English label)
	ref: str | None = None  # a lookup name, see LOOKUPS
	example: str = ""
	help: str = ""


@dataclass
class Spec:
	name: str
	title: str  # English source text
	columns: tuple[Col, ...]
	save: Callable[[resources.Context, dict[str, Any]], Any]
	prepare: Callable[[resources.Context, dict[str, Any], Session], dict[str, Any]] | None = None
	note: str = ""  # shown in the guide sheet (English source text)


# what an "existing thing" column is matched on, and the model it comes from
LOOKUPS: dict[str, tuple[type, str]] = {
	"employees": (Employee, "code"),
	"work-locations": (WorkLocation, "name"),
	"provinces": (Province, "name"),
	"departments": (Department, "name"),
	"designations": (Designation, "name"),
	"shifts": (Shift, "name"),
	"work-patterns": (WorkPattern, "name"),
	"approval-processes": (ApprovalProcess, "name"),
	"leave-types": (LeaveType, "name"),
	"devices": (AttendanceDevice, "name"),
}


def _options(session: Session, ref: str) -> dict[str, int]:
	model, attribute = LOOKUPS[ref]
	column = getattr(model, attribute)
	return {str(v).strip(): i for i, v in session.execute(select(model.id, column)) if v is not None}


# the specs -----------------------------------------------------------------------------------------

REQUEST_TYPES = (
	("daily_leave", _("Daily leave")),
	("hourly_leave", _("Hourly leave")),
	("mission", _("Mission")),
	("overtime", _("Overtime")),
	("attendance_correction", _("Manual time entry")),
)
TARGETS = (
	("employee", _("One employee")),
	("designation", _("A designation")),
	("department", _("A department")),
	("company", _("The whole company")),
)


def _generic(name: str) -> Callable[[resources.Context, dict[str, Any]], Any]:
	def save(ctx: resources.Context, values: dict[str, Any]) -> Any:
		return resources.save(ctx, resources.resource(name), values)

	return save


def _days(ctx: resources.Context, values: dict[str, Any], session: Session) -> dict[str, Any]:
	"""A work pattern row: seven weekday columns, or a rotation like ``Morning*14, Night*14``."""
	shifts = _options(session, "shifts")

	def shift_id(name: str) -> int:
		if name not in shifts:
			raise ValidationError(_("The shift {0} does not exist"), name)
		return shifts[name]

	days: list[dict[str, Any]] = []
	if values.get("pattern_type") == "rotating":
		cycle = 1
		for part in re.split(r"[,،\n]+", str(values.pop("rotation", "") or "")):
			part = part.strip()
			if not part:
				continue
			name, _star, count = part.rpartition("*")
			if not _star:
				name, count = part, "1"
			if not count.strip().isdigit() or int(count) < 1:
				raise ValidationError(_("The rotation {0} must look like Morning*14"), part)
			for _n in range(int(count)):
				days.append({"cycle_day": cycle, "shift_id": shift_id(name.strip())})
				cycle += 1
		if not days:
			raise ValidationError(_("Enter the rotation, for example Morning*14, Night*14"))
		for i in range(7):
			values.pop(f"weekday_{i}", None)
	else:
		values.pop("rotation", None)
		for i in range(7):
			name = values.pop(f"weekday_{i}", None)
			days.append({"weekday": i, "shift_id": shift_id(name) if name else None})
	return {**values, "days": days}


def _steps(ctx: resources.Context, values: dict[str, Any], session: Session) -> dict[str, Any]:
	"""An approval process row: up to four steps, each ``parent``, ``ceo`` or an employee code."""
	people = _options(session, "employees")
	steps = []
	for i in range(1, 5):
		cell = values.pop(f"step_{i}", None)
		if not cell:
			continue
		if cell in ("parent", "ceo"):
			steps.append({"position": len(steps) + 1, "approver_type": cell})
		elif cell in people:
			steps.append(
				{"position": len(steps) + 1, "approver_type": "specific_person", "approver_id": people[cell]}
			)
		else:
			raise ValidationError(_("Step {0}: {1} is not parent, CEO or an employee code"), i, cell)
	return {**values, "steps": steps}


def _parameters(ctx: resources.Context, values: dict[str, Any], session: Session) -> dict[str, Any]:
	law = baseline_rule_set()
	parameters = {p.name: values.pop(p.name, None) for p in PARAMETERS}
	return {
		**values,
		"parameters": {k: (v if v is not None else getattr(law, k)) for k, v in parameters.items()},
	}


def _report_columns(ctx: resources.Context, values: dict[str, Any], session: Session) -> dict[str, Any]:
	keys = [k.strip() for k in re.split(r"[,،\n]+", str(values.pop("column_keys", "") or "")) if k.strip()]
	return {**values, "columns": [{"field": k} for k in keys]}


def _holiday(ctx: resources.Context, values: dict[str, Any]) -> Any:
	return calendar_service.create_holiday(ctx.session, values, ctx.actor.user_id)


def _device(ctx: resources.Context, values: dict[str, Any]) -> Any:
	return devices.save_device(ctx.session, values, ctx.actor.user_id)


def _device_user(ctx: resources.Context, values: dict[str, Any]) -> Any:
	return devices.save_mapping(ctx.session, values, ctx.actor.user_id)


def _employee(ctx: resources.Context, values: dict[str, Any]) -> Any:
	return employees.save(ctx, values)


def _c(key, label, kind=TEXT, required=False, choices=(), ref=None, example="", help=""):
	return Col(key, label, kind, required, tuple(choices), ref, example, help)


GENDERS = (("male", _("Male")), ("female", _("Female")))
STATUSES = (("active", _("Active")), ("suspended", _("Suspended")), ("left", _("Left")))
CONTRACTS = (("full_time", _("Full time")), ("part_time", _("Part time")), ("hourly", _("Hourly")))
HOLIDAY_TYPES = (
	("official_solar", _("Official (solar)")),
	("official_lunar", _("Official (lunar)")),
	("official_occasional", _("Official (announced)")),
	("office_closure", _("Office closure")),
	("company_internal", _("Company holiday")),
)
SCOPES = (
	("country", _("Whole country")),
	("province", _("One province")),
	("work_location", _("One work location")),
	("company", _("Whole company")),
)
CLOSURES = (
	("paid_holiday", _("Paid holiday")),
	("employer_paid", _("Paid by the employer, counted as work")),
	("workday", _("Normal workday (make up later)")),
)
HOLIDAY_STATUSES = (("confirmed", _("Confirmed")), ("predicted", _("Predicted")))
LEAVE_CATEGORIES = (
	("earned", _("Earned")),
	("sick", _("Sick")),
	("unpaid", _("Unpaid")),
	("marriage_bereavement", _("Marriage or bereavement")),
	("hajj", _("Hajj")),
	("maternity", _("Maternity")),
	("paternity", _("Paternity")),
	("other", _("Other")),
)
DIRECTION_MODES = (
	("from_device", _("From Device")),
	("always_in", _("Always In")),
	("always_out", _("Always Out")),
	("pair_by_order", _("Pair by order")),
)
SOURCES = (("summary", _("One row per person")), ("daily", _("One row per day")))
TIME_UNITS = (
	("hours_minutes", _("Hours:minutes")),
	("decimal_hours", _("Decimal hours")),
	("days", _("Days")),
)
EXPORTS = (("excel", _("Excel")), ("pdf", _("PDF")), ("csv", _("Text (CSV)")))

SPECS: dict[str, Spec] = {
	s.name: s
	for s in (
		Spec(
			"work-locations",
			_("Work locations"),
			(
				_c("name", _("Name"), required=True, example="دفتر مرکزی"),
				_c("province_id", _("Province"), REF, True, ref="provinces", example="تهران"),
				_c("city", _("City"), example="تهران"),
				_c("address", _("Address")),
				_c("is_active", _("Active"), BOOL, example="بله"),
			),
			_generic("work-locations"),
		),
		Spec(
			"departments",
			_("Departments"),
			(_c("name", _("Name"), required=True, example="فروش"),),
			_generic("departments"),
		),
		Spec(
			"designations",
			_("Designations"),
			(
				_c("name", _("Name"), required=True, example="مدیر فروش"),
				_c(
					"parent_id",
					_("Parent designation"),
					REF,
					ref="designations",
					help=_("Put parents above their children in the sheet"),
				),
				_c("hard_and_harmful", _("Hard and harmful work"), BOOL),
				_c("committee_ruling_no", _("Committee ruling number")),
				_c("committee_ruling_date", _("Committee ruling date"), JDATE),
				_c("description", _("Description")),
			),
			_generic("designations"),
			note=_("A designation's parent must already exist or be on an earlier row."),
		),
		Spec(
			"designation-assignments",
			_("Designations held"),
			(
				_c(
					"employee_id",
					_("Employee"),
					REF,
					True,
					ref="employees",
					example="1001",
					help=_("The personnel code"),
				),
				_c("designation_id", _("Designation"), REF, True, ref="designations", example="کارشناس فروش"),
				_c("valid_from", _("From"), JDATE, True, example="1405/01/01"),
				_c("valid_to", _("Until"), JDATE),
				_c("note", _("Note")),
			),
			_generic("designation-assignments"),
		),
		Spec(
			"employees",
			_("Employees"),
			(
				_c("code", _("Personnel code"), required=True, example="1001"),
				_c(
					"detail_code",
					_("Detail code"),
					help=_("The accounting detail code (کد تفصیلی)"),
					example="10020030",
				),
				_c("first_name", _("First name"), required=True, example="سارا"),
				_c("last_name", _("Last name"), required=True, example="کریمی"),
				_c("national_id", _("National code"), example="0012345678"),
				_c("gender", _("Gender"), CHOICE, choices=GENDERS),
				_c("mobile", _("Mobile"), example="09121234567"),
				_c("email", _("Email")),
				_c("start_date", _("Start date"), JDATE, True, example="1405/01/01"),
				_c("end_date", _("End date"), JDATE),
				_c("status", _("Status"), CHOICE, choices=STATUSES, example=_("Active")),
				_c("contract_type", _("Contract"), CHOICE, choices=CONTRACTS),
				_c("work_location_id", _("Work location"), REF, ref="work-locations"),
				_c("department_id", _("Department"), REF, ref="departments"),
				_c(
					"designation_id",
					_("Designation"),
					REF,
					ref="designations",
					help=_("The first designation they hold"),
				),
				_c("designation_from", _("Designation from"), JDATE, help=_("Empty: the start date")),
				_c("birth_date", _("Birth date"), JDATE, help=_("Only for young-worker rules")),
				_c("overtime_consent_on_file", _("Written overtime consent on file"), BOOL),
				_c("overtime_consent_date", _("Overtime consent date"), JDATE),
				_c("biometric_consent_date", _("Biometric consent date"), JDATE),
				_c("uses_card_instead_of_biometrics", _("Uses a card instead of biometrics"), BOOL),
				_c("nursing_until", _("Nursing breaks until"), JDATE),
				_c("username", _("Username"), help=_("Empty: the personnel code")),
				_c(
					"password",
					_("Password"),
					required=True,
					help=_("At least 8 characters; the person can change it"),
				),
				_c("is_admin", _("System admin"), BOOL),
			),
			_employee,
			note=_("Every employee is also a user: the username and password create the login."),
		),
		Spec(
			"shifts",
			_("Shifts"),
			(
				_c("name", _("Name"), required=True, example="اداری"),
				_c(
					"shift_type",
					_("Type"),
					CHOICE,
					True,
					choices=(("fixed", _("Fixed")), ("flexible", _("Flexible"))),
					example=_("Fixed"),
				),
				_c("start_time", _("Start"), TIME, True, example="08:00"),
				_c(
					"end_time",
					_("End"),
					TIME,
					True,
					example="16:30",
					help=_("Earlier than the start: the shift ends the next day"),
				),
				_c("latest_start_time", _("Latest start"), TIME),
				_c("required_minutes", _("Required minutes"), NUMBER),
				_c("break_start_time", _("Break start"), TIME, example="12:30"),
				_c("break_end_time", _("Break end"), TIME, example="13:10"),
				_c("break_minutes", _("Break minutes without punching"), NUMBER),
				_c("grace_late_minutes", _("Grace for late arrival (minutes)"), NUMBER),
				_c("grace_early_minutes", _("Grace for early leave (minutes)"), NUMBER),
				_c("overtime_after_minutes", _("Extra work counts after (minutes)"), NUMBER),
				_c("color", _("Colour")),
			),
			_generic("shifts"),
		),
		Spec(
			"work-patterns",
			_("Work patterns"),
			(
				_c("name", _("Name"), required=True, example="هفتگی اداری"),
				_c(
					"pattern_type",
					_("Type"),
					CHOICE,
					True,
					choices=(("weekly", _("Weekly")), ("rotating", _("Rotating"))),
					example=_("Weekly"),
				),
				_c(
					"weekly_rest_day",
					_("Weekly rest day"),
					NUMBER,
					help=_("0 = Saturday ... 6 = Friday; empty: Friday"),
				),
				_c("cycle_start_date", _("First day of the cycle"), JDATE, help=_("Rotating patterns only")),
				*[
					_c(
						f"weekday_{i}",
						WEEKDAY_NAMES[i],
						REF,
						ref="shifts",
						help=_("Weekly patterns: the shift, empty = day off"),
					)
					for i in range(7)
				],
				_c(
					"rotation",
					_("Rotation"),
					help=_("Rotating patterns: shift*days, e.g. Morning*14, Night*14"),
				),
				_c("description", _("Description")),
				_c("change_reason", _("Reason for the change")),
			),
			_generic("work-patterns"),
			prepare=_days,
			note=_(
				"A weekly pattern fills the seven weekday columns; a rotating pattern fills the rotation column."
			),
		),
		Spec(
			"work-pattern-assignments",
			_("Pattern assignments"),
			(
				_c("target_type", _("Applies to"), CHOICE, True, choices=TARGETS, example=_("One employee")),
				_c(
					"employee_id",
					_("Employee"),
					REF,
					ref="employees",
					help=_("When it applies to one employee"),
				),
				_c(
					"designation_id",
					_("Designation"),
					REF,
					ref="designations",
					help=_("When it applies to a designation"),
				),
				_c(
					"department_id",
					_("Department"),
					REF,
					ref="departments",
					help=_("When it applies to a department"),
				),
				_c("work_pattern_id", _("Work pattern"), REF, True, ref="work-patterns"),
				_c(
					"cycle_offset_days",
					_("Start offset in the cycle (days)"),
					NUMBER,
					help=_("Rotating patterns: shifts this group along the cycle"),
				),
				_c("valid_from", _("From"), JDATE, True, example="1405/01/01"),
				_c("valid_to", _("Until"), JDATE),
				_c("note", _("Note")),
			),
			_generic("work-pattern-assignments"),
		),
		Spec(
			"shift-overrides",
			_("One-day shift changes"),
			(
				_c("employee_id", _("Employee"), REF, True, ref="employees"),
				_c("jdate", _("Date"), JDATE, True, example="1405/07/20"),
				_c("shift_id", _("Shift (empty = day off)"), REF, ref="shifts"),
				_c("reason", _("Reason"), required=True),
			),
			_generic("shift-overrides"),
		),
		Spec(
			"leave-types",
			_("Leave types"),
			(
				_c("name", _("Name"), required=True),
				_c("category", _("Category"), CHOICE, True, choices=LEAVE_CATEGORIES, example=_("Other")),
				_c("is_paid", _("Paid"), BOOL),
				_c("deducts_from_balance", _("Deducted from the leave balance"), BOOL),
				_c("allow_hourly", _("Can be hourly"), BOOL),
				_c("max_days_per_request", _("Most days in one request"), NUMBER),
				_c("once_in_service", _("Once in the whole service"), BOOL),
				_c("rule_limit", _("Limit from the rule set (parameter)")),
				_c("description", _("Description")),
			),
			_generic("leave-types"),
		),
		Spec(
			"approval-processes",
			_("Approval processes"),
			(
				_c("name", _("Name"), required=True),
				_c("description", _("Description")),
				_c("step_1", _("Step 1"), help=_("parent, ceo, or the personnel code of a specific person")),
				_c("step_2", _("Step 2"), help=_("parent, ceo, or the personnel code of a specific person")),
				_c("step_3", _("Step 3"), help=_("parent, ceo, or the personnel code of a specific person")),
				_c("step_4", _("Step 4"), help=_("parent, ceo, or the personnel code of a specific person")),
			),
			_generic("approval-processes"),
			prepare=_steps,
		),
		Spec(
			"approval-overrides",
			_("Approval exceptions"),
			(
				_c("employee_id", _("Employee"), REF, True, ref="employees"),
				_c(
					"request_type",
					_("Request type"),
					CHOICE,
					True,
					choices=REQUEST_TYPES,
					example=_("Overtime"),
				),
				_c("process_id", _("Approval process"), REF, True, ref="approval-processes"),
			),
			_generic("approval-overrides"),
		),
		Spec(
			"rule-sets",
			_("Labor rule sets"),
			(
				_c("name", _("Name"), required=True),
				_c("effective_from", _("Effective from"), JDATE, True, example="1406/01/01"),
				_c("description", _("Description")),
				_c("change_reason", _("Reason for the change")),
				*[
					_c(p.name, p.name, NUMBER, help=f"Article {p.article}; empty = the legal value")
					for p in PARAMETERS
				],
			),
			_generic("rule-sets"),
			prepare=_parameters,
			note=_("A value may be better for the worker than the law, never worse."),
		),
		Spec(
			"report-templates",
			_("Report templates"),
			(
				_c("name", _("Name"), required=True),
				_c("source", _("Rows"), CHOICE, True, choices=SOURCES, example=_("One row per person")),
				_c("time_unit", _("Show times as"), CHOICE, choices=TIME_UNITS),
				_c("export_format", _("Format"), CHOICE, choices=EXPORTS),
				_c(
					"column_keys",
					_("Columns"),
					required=True,
					help=_("Field keys separated by commas, e.g. employee, employee_name, worked_minutes"),
				),
				_c("description", _("Description")),
			),
			_generic("report-templates"),
			prepare=_report_columns,
		),
		Spec(
			"holidays",
			_("Holidays"),
			(
				_c("jdate", _("Date"), JDATE, True, example="1405/07/20"),
				_c("title", _("Title"), required=True, example="تعطیلی شرکت"),
				_c(
					"holiday_type",
					_("Type"),
					CHOICE,
					True,
					choices=HOLIDAY_TYPES,
					example=_("Company holiday"),
				),
				_c("scope", _("Applies to"), CHOICE, choices=SCOPES, example=_("Whole company")),
				_c("province_id", _("Province"), REF, ref="provinces"),
				_c("work_location_id", _("Work location"), REF, ref="work-locations"),
				_c("closure_pay_mode", _("How is it paid?"), CHOICE, choices=CLOSURES),
				_c("status", _("Status"), CHOICE, choices=HOLIDAY_STATUSES),
				_c("note", _("Note")),
			),
			_holiday,
		),
		Spec(
			"devices",
			_("Attendance devices"),
			(
				_c("name", _("Name"), required=True, example="VF600 ورودی"),
				_c("adapter_type", _("Device Type"), required=True, example="zkteco"),
				_c("work_location_id", _("Work location"), REF, ref="work-locations"),
				_c("host", _("IP Address"), example="192.168.1.201"),
				_c("port", _("Port"), NUMBER, example="4370"),
				_c("comm_key", _("Comm key")),
				_c("direction_mode", _("Punch Direction"), CHOICE, choices=DIRECTION_MODES),
				_c("polling_interval_minutes", _("Read Every (Minutes)"), NUMBER),
				_c("timeout_seconds", _("Timeout (Seconds)"), NUMBER),
				_c("use_udp", _("Use UDP"), BOOL),
				_c("enabled", _("Enabled"), BOOL),
			),
			_device,
		),
		Spec(
			"device-users",
			_("Device users"),
			(
				_c("device_id", _("Device"), REF, True, ref="devices"),
				_c("device_user_id", _("Device user id"), required=True, example="12"),
				_c("device_user_name", _("Name on the device")),
				_c("employee_id", _("Employee"), REF, ref="employees"),
				_c("valid_from", _("From"), JDATE),
				_c("valid_to", _("Until"), JDATE),
			),
			_device_user,
		),
	)
}


def spec(name: str) -> Spec:
	try:
		return SPECS[name]
	except KeyError:
		raise NotFound(_("Not found")) from None


def available() -> list[dict[str, str]]:
	return [{"name": s.name, "title": s.title} for s in SPECS.values()]


# the sample workbook --------------------------------------------------------------------------------


def _choice_label(choice: tuple[str, str], language: str) -> str:
	return translate(choice[1], language)


def _format_of(col: Col, language: str) -> str:
	return {
		TEXT: translate(_("Text"), language),
		NUMBER: translate(_("Number"), language),
		BOOL: translate(_("Yes or no"), language),
		JDATE: translate(_("Jalali date, e.g. 1405/07/14"), language),
		TIME: translate(_("Time, e.g. 08:30"), language),
		CHOICE: translate(_("Choose from the list"), language),
		REF: translate(_("Choose from the list"), language),
	}[col.kind]


def template(session: Session, name: str, language: str = "fa") -> bytes:
	"""The sample workbook for one import."""
	s = spec(name)
	t = lambda text: translate(text, language)  # noqa: E731
	wb = Workbook()
	data = wb.active
	data.title = t(_("Data entry"))[:31]
	guide = wb.create_sheet(t(_("Guide"))[:31])
	lists = wb.create_sheet("lists")
	lists.sheet_state = "hidden"
	for sheet in (data, guide):
		sheet.sheet_view.rightToLeft = True
	header_fill = PatternFill("solid", fgColor="0F766E")
	required_fill = PatternFill("solid", fgColor="B45309")

	list_column = 0
	for index, col in enumerate(s.columns, start=1):
		letter = get_column_letter(index)
		cell = data.cell(row=1, column=index, value=t(col.label) + (" *" if col.required else ""))
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = required_fill if col.required else header_fill
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
		data.column_dimensions[letter].width = max(14, min(34, len(str(cell.value)) + 6))
		span = f"{letter}2:{letter}{FORMATTED_ROWS + 1}"
		options: list[str] = []
		if col.kind == CHOICE:
			options = [_choice_label(c, language) for c in col.choices]
		elif col.kind == REF and col.ref:
			options = sorted(_options(session, col.ref))
		elif col.kind == BOOL:
			options = [t(_("Yes")), t(_("No"))]
		if options:
			list_column += 1
			for i, option in enumerate(options, start=1):
				lists.cell(row=i, column=list_column, value=option)
			ref = (
				f"lists!${get_column_letter(list_column)}$1:${get_column_letter(list_column)}${len(options)}"
			)
			validation = DataValidation(
				type="list", formula1=ref, allow_blank=not col.required, showErrorMessage=True
			)
			validation.error = t(_("Choose a value from the list"))
			data.add_data_validation(validation)
			validation.add(span)
		elif col.kind == NUMBER:
			validation = DataValidation(
				type="decimal",
				operator="greaterThanOrEqual",
				formula1="0",
				allow_blank=True,
				showErrorMessage=True,
			)
			validation.error = t(_("Enter a number"))
			data.add_data_validation(validation)
			validation.add(span)
		elif col.kind == JDATE:
			validation = DataValidation(
				type="custom",
				formula1=f'=OR({letter}2="",AND(LEN({letter}2)=10,MID({letter}2,5,1)="/",MID({letter}2,8,1)="/"))',
				allow_blank=True,
				showErrorMessage=True,
			)
			validation.error = t(_("Enter the date as 1405/07/14"))
			data.add_data_validation(validation)
			validation.add(span)
		elif col.kind == TIME:
			validation = DataValidation(
				type="custom",
				formula1=f'=OR({letter}2="",AND(LEN({letter}2)=5,MID({letter}2,3,1)=":"))',
				allow_blank=True,
				showErrorMessage=True,
			)
			validation.error = t(_("Enter the time as 08:30"))
			data.add_data_validation(validation)
			validation.add(span)
		# text-like columns stay text, so nothing is turned into a date or a number by Excel
		fmt = "General" if col.kind == NUMBER else "@"
		for row in range(2, FORMATTED_ROWS + 2):
			data.cell(row=row, column=index).number_format = fmt
	data.row_dimensions[1].height = 34
	data.freeze_panes = "A2"

	guide["A1"] = t(s.title)
	guide["A1"].font = Font(bold=True, size=14)
	guide["A2"] = t(
		"Fill the first sheet from row 2. Columns marked * are required. Do not change the header row."
	)
	guide["A3"] = t(s.note) if s.note else ""
	headers = [
		t(_("Column")),
		t(_("Required")),
		t(_("Format")),
		t(_("Allowed values")),
		t(_("Example")),
		t(_("Note")),
	]
	for i, h in enumerate(headers, start=1):
		cell = guide.cell(row=5, column=i, value=h)
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = header_fill
	for r, col in enumerate(s.columns, start=6):
		allowed = ""
		if col.kind == CHOICE:
			allowed = "، ".join(_choice_label(c, language) for c in col.choices)
		elif col.kind == BOOL:
			allowed = f"{t('Yes')}، {t('No')}"
		values = [
			t(col.label),
			t(_("Yes")) if col.required else "",
			_format_of(col, language),
			allowed,
			t(col.example) if col.example else "",
			t(col.help) if col.help else "",
		]
		for i, v in enumerate(values, start=1):
			guide.cell(row=r, column=i, value=v).alignment = Alignment(wrap_text=True, vertical="top")
	for letter, width in zip("ABCDEF", (28, 10, 28, 40, 18, 46), strict=True):
		guide.column_dimensions[letter].width = width
	out = io.BytesIO()
	wb.save(out)
	return out.getvalue()


# reading it back --------------------------------------------------------------------------------------


@dataclass
class RowError:
	row: int
	message: str


@dataclass
class Result:
	created: int = 0
	errors: list[RowError] = field(default_factory=list)

	@property
	def imported(self) -> bool:
		return not self.errors and self.created > 0

	def as_dict(self) -> dict[str, Any]:
		return {
			"imported": self.imported,
			"created": self.created,
			"errors": [{"row": e.row, "message": e.message} for e in self.errors],
		}


def _text(value: Any) -> str:
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		return str(int(value))
	if isinstance(value, datetime | date | time):
		raise ValidationError(_("This cell is a real Excel date or time; type it as text, e.g. 1405/07/14"))
	return to_ascii_digits(str(value)).strip()


def _convert(col: Col, raw: Any, session: Session, language: str, cache: dict[str, dict[str, int]]) -> Any:
	if col.kind == TIME and isinstance(raw, time):
		raw = raw.strftime("%H:%M")
	elif col.kind == TIME and isinstance(raw, float) and 0 <= raw < 1:  # a time typed as an Excel time
		minutes = round(raw * 1440)
		raw = f"{minutes // 60:02d}:{minutes % 60:02d}"
	text = _text(raw)
	if text == "":
		return None
	if col.kind == TEXT:
		return text
	if col.kind == NUMBER:
		try:
			number = Decimal(text.replace(",", "").replace("٫", "."))
		except Exception:
			raise ValidationError(_("{0} must be a number"), translate(col.label, language)) from None
		return int(number) if number == number.to_integral_value() else float(number)
	if col.kind == BOOL:
		word = text.lower()
		if word in YES_WORDS or word == translate(_("Yes"), language).lower():
			return True
		if word in NO_WORDS or word == translate(_("No"), language).lower():
			return False
		raise ValidationError(_("{0}: enter Yes or No"), translate(col.label, language))
	if col.kind == JDATE:
		try:
			return normalize_jdate(text.replace("-", "/"))
		except InvalidJalaliDate:
			raise ValidationError(
				_("{0}: {1} is not a valid Jalali date"), translate(col.label, language), text
			) from None
	if col.kind == TIME:
		clock = normalize_time(text)
		if clock is None:
			raise ValidationError(_("{0}: enter the time as HH:MM"), translate(col.label, language))
		return clock[:5]
	if col.kind == CHOICE:
		for code, label in col.choices:
			if text in (code, label, translate(label, language)):
				return code
		raise ValidationError(
			_("{0}: {1} is not one of the allowed values"), translate(col.label, language), text
		)
	if col.kind == REF:
		if col.ref not in cache:
			cache[col.ref] = _options(session, col.ref)
		if text not in cache[col.ref]:
			raise ValidationError(_("{0}: {1} does not exist"), translate(col.label, language), text)
		return cache[col.ref][text]
	return text


def _header_map(sheet, s: Spec, language: str) -> dict[int, Col]:
	by_name: dict[str, Col] = {}
	for col in s.columns:
		for name in (translate(col.label, language), col.label, col.key):
			by_name[name.strip().lower()] = col
	mapping: dict[int, Col] = {}
	for index, cell in enumerate(sheet[1], start=1):
		label = _text(cell.value).removesuffix("*").strip().lower()
		if label in by_name:
			mapping[index] = by_name[label]
	missing = [translate(c.label, language) for c in s.columns if c.required and c not in mapping.values()]
	if missing:
		raise ValidationError(_("The sheet has no column {0}"), "، ".join(missing))
	return mapping


def run(ctx: resources.Context, name: str, content: bytes) -> Result:
	"""Import a filled workbook. Nothing is saved unless every row is valid."""
	s = spec(name)
	language = ctx.actor.language
	try:
		workbook = load_workbook(io.BytesIO(content), data_only=True)
	except Exception:
		raise ValidationError(_("The file is not an Excel workbook (.xlsx)")) from None
	sheet = workbook.worksheets[0]
	mapping = _header_map(sheet, s, language)
	result = Result()
	cache: dict[str, dict[str, int]] = {}
	outer = ctx.session.begin_nested()
	try:
		for row_number, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
			if not any(c not in (None, "") for c in cells):
				continue
			if result.created + len(result.errors) >= MAX_ROWS:
				raise ValidationError(_("The sheet has more than {0} rows"), MAX_ROWS)
			inner = ctx.session.begin_nested()
			try:
				values: dict[str, Any] = {}
				problems: list[str] = []
				for index, col in mapping.items():
					raw = cells[index - 1] if index - 1 < len(cells) else None
					try:
						value = _convert(col, raw, ctx.session, language, cache)
					except HozurError as e:
						problems.append(e.translated(language))
						continue
					if value is None and col.required:
						problems.append(
							translate(_("{0} is required"), language).format(translate(col.label, language))
						)
					if value is not None:
						values[col.key] = value
				if problems:
					raise ValidationError("; ".join(problems).replace("{", "{{").replace("}", "}}"))
				if s.prepare:
					values = s.prepare(ctx, values, ctx.session)
				s.save(ctx, values)
				inner.commit()
				cache.pop(s.name, None)  # what this row created can be used by the rows below it
				result.created += 1
			except HozurError as e:
				inner.rollback()
				result.errors.append(RowError(row_number, e.translated(language)))
	except Exception:
		outer.rollback()
		raise
	if result.errors:
		outer.rollback()
		result.created = 0
	else:
		outer.commit()
	return result
