"""Report exporters: Excel, text (CSV / delimited / fixed width) and PDF. Pure Python.

columns: [{"key", "label", "kind", "width"}]; rows: [{key: raw value}] where minutes are ints.
"""

from __future__ import annotations

import csv
import hashlib
import html
import io
import json
from collections.abc import Sequence

from hozur.jalali.dates import to_ascii_digits, to_persian_digits
from hozur.reports.catalog import COUNT, DATE, DAYS, MINUTES

HOURS_MINUTES, DECIMAL_HOURS, DAYS_UNIT = "Hours:Minutes", "Decimal Hours", "Days"
DELIMITERS = {"Comma": ",", "Tab": "\t", "Pipe": "|", "Semicolon": ";"}
FIXED_WIDTH = "Fixed Width"
# Windows-1256 has no Persian yeh: older payroll software expects the Arabic one.
_CP1256 = str.maketrans({"ی": "ي", "ى": "ي", "ۀ": "ه"})


def hours_minutes(minutes: int) -> str:
	sign = "-" if minutes < 0 else ""
	minutes = abs(int(minutes))
	return f"{sign}{minutes // 60}:{minutes % 60:02d}"


def display_value(value, kind: str, unit: str, day_minutes: int):
	"""The value as shown (text, PDF). Minutes follow the report's time unit."""
	if value is None:
		return ""
	if kind == MINUTES:
		if unit == DECIMAL_HOURS:
			return f"{value / 60:.2f}"
		if unit == DAYS_UNIT:
			return f"{value / day_minutes:.2f}"
		return hours_minutes(value)
	if kind == DAYS:
		return f"{float(value):.2f}".rstrip("0").rstrip(".") if value else "0"
	if kind == COUNT:
		return str(int(value))
	return str(value)


def totals(columns: Sequence[dict], rows: Sequence[dict]) -> dict:
	return {
		c["key"]: sum(r.get(c["key"]) or 0 for r in rows)
		for c in columns
		if c["kind"] in (MINUTES, DAYS, COUNT)
	}


def content_hash(columns: Sequence[dict], rows: Sequence[dict]) -> str:
	"""SHA-256 of the report data, printed on the PDF so a tampered copy can be detected."""
	payload = json.dumps(
		{"columns": [c["key"] for c in columns], "rows": [[r.get(c["key"]) for c in columns] for r in rows]},
		ensure_ascii=False,
		sort_keys=True,
		default=str,
	)
	return hashlib.sha256(payload.encode("utf-8")).hexdigest()


# Excel ------------------------------------------------------------------------------------


def to_xlsx(
	title: str,
	meta: Sequence[str],
	columns: Sequence[dict],
	rows: Sequence[dict],
	unit: str,
	day_minutes: int,
) -> bytes:
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill
	from openpyxl.utils import get_column_letter

	wb = Workbook()
	ws = wb.active
	ws.title = (title or "Report")[:31]
	ws.sheet_view.rightToLeft = True
	font, bold = Font(name="Tahoma", size=10), Font(name="Tahoma", size=10, bold=True)
	ws.append([title])
	ws["A1"].font = Font(name="Tahoma", size=12, bold=True)
	for line in meta:
		ws.append([line])
	ws.append([])
	header_row = ws.max_row + 1
	ws.append([c["label"] for c in columns])
	for cell in ws[header_row]:
		cell.font, cell.fill = bold, PatternFill("solid", fgColor="EEEEEE")
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

	def write_row(values: dict, is_total=False):
		ws.append([None] * len(columns))
		r = ws.max_row
		for i, c in enumerate(columns, start=1):
			value, cell = values.get(c["key"]), ws.cell(row=r, column=i)
			cell.font = bold if is_total else font
			if value is None or value == "":
				continue
			if c["kind"] == MINUTES:
				if unit == DECIMAL_HOURS:
					cell.value, cell.number_format = round(value / 60, 2), "0.00"
				elif unit == DAYS_UNIT:
					cell.value, cell.number_format = round(value / day_minutes, 2), "0.00"
				else:
					cell.value, cell.number_format = value / 1440, "[h]:mm"
			elif c["kind"] in (DAYS, COUNT):
				cell.value = round(float(value), 4) if c["kind"] == DAYS else int(value)
			else:
				cell.value = str(value)
				if c["kind"] == DATE:
					cell.alignment = Alignment(horizontal="center")

	for row in rows:
		write_row(row)
	if rows and any(c["kind"] in (MINUTES, DAYS, COUNT) for c in columns):
		write_row({**totals(columns, rows), columns[0]["key"]: "جمع"}, is_total=True)
	for i, c in enumerate(columns, start=1):
		ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(c["label"]) + 4))
	ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
	buf = io.BytesIO()
	wb.save(buf)
	return buf.getvalue()


# text -------------------------------------------------------------------------------------


def to_text(
	columns: Sequence[dict],
	rows: Sequence[dict],
	unit: str,
	day_minutes: int,
	delimiter: str = "Comma",
	encoding: str = "UTF-8",
	header: bool = True,
) -> bytes:
	"""Delimited or fixed-width text with ASCII digits, for payroll software."""
	lines = []
	if header:
		lines.append([c["label"] for c in columns])
	for row in rows:
		lines.append(
			[to_ascii_digits(display_value(row.get(c["key"]), c["kind"], unit, day_minutes)) for c in columns]
		)
	if delimiter == FIXED_WIDTH:
		widths = [int(c.get("width") or 12) for c in columns]
		text = (
			"\r\n".join(
				"".join(str(v)[:w].ljust(w) for v, w in zip(line, widths, strict=False)) for line in lines
			)
			+ "\r\n"
		)
	else:
		buf = io.StringIO()
		csv.writer(buf, delimiter=DELIMITERS.get(delimiter, ","), lineterminator="\r\n").writerows(lines)
		text = buf.getvalue()
	if encoding == "Windows-1256":
		return text.translate(_CP1256).encode("cp1256", errors="replace")
	if encoding == "UTF-8 with BOM":
		return text.encode("utf-8-sig")
	return text.encode("utf-8")


# PDF --------------------------------------------------------------------------------------

PDF_CSS = """
@font-face {{ font-family: Vazirmatn; src: url({font}/Vazirmatn-Regular.ttf); font-weight: 400; }}
@font-face {{ font-family: Vazirmatn; src: url({font}/Vazirmatn-Bold.ttf); font-weight: 700; }}
@page {{
	size: A4 {orientation}; margin: 14mm 10mm 18mm;
	@bottom-left {{ content: "صفحه " counter(page, persian) " از " counter(pages, persian); font: 8pt Vazirmatn; }}
	@bottom-right {{ content: "{footer}"; font: 7pt Vazirmatn; color: #555; }}
}}
body {{ font-family: Vazirmatn; font-size: 8.5pt; color: #111; }}
h1 {{ font-size: 13pt; margin: 0 0 1.5mm; }}
.meta {{ color: #444; margin: 0 0 4mm; line-height: 1.6; }}
table {{ border-collapse: collapse; width: 100%; }}
th, td {{ border: 0.4pt solid #999; padding: 1.2mm 1.8mm; text-align: right; vertical-align: top; }}
th {{ background: #eee; font-weight: 700; }}
td.n {{ text-align: left; direction: ltr; white-space: nowrap; }}
tr.total td {{ font-weight: 700; background: #f6f6f6; }}
.hash {{ margin-top: 4mm; font-size: 7pt; color: #555; }}
.hash span {{ direction: ltr; unicode-bidi: embed; font-family: monospace; }}
"""


def to_pdf(
	title: str,
	meta: Sequence[str],
	columns: Sequence[dict],
	rows: Sequence[dict],
	unit: str,
	day_minutes: int,
	font_url: str,
) -> bytes:
	from weasyprint import HTML

	esc = html.escape
	digest = content_hash(columns, rows)
	numeric = {c["key"] for c in columns if c["kind"] in (MINUTES, DAYS, COUNT)}

	def td(c, value):
		text = to_persian_digits(display_value(value, c["kind"], unit, day_minutes))
		return (
			f'<td class="n">{esc(text)}</td>'
			if c["key"] in numeric
			else f"<td>{esc(to_persian_digits(text))}</td>"
		)

	body = "".join("<tr>" + "".join(td(c, r.get(c["key"])) for c in columns) + "</tr>" for r in rows)
	if rows and numeric:
		sums = totals(columns, rows)
		body += (
			'<tr class="total">'
			+ "".join(
				td(c, sums.get(c["key"]))
				if c["key"] in numeric
				else (f"<td>{esc('جمع')}</td>" if i == 0 else "<td></td>")
				for i, c in enumerate(columns)
			)
			+ "</tr>"
		)
	css = PDF_CSS.format(
		font=font_url,
		orientation="landscape" if len(columns) > 7 else "portrait",
		footer=esc(title).replace('"', ""),
	)
	document = f"""<!doctype html><html lang="fa" dir="rtl"><head><meta charset="utf-8"><style>{css}</style></head><body>
<h1>{esc(title)}</h1>
<div class="meta">{"<br>".join(esc(to_persian_digits(m)) for m in meta)}</div>
<table><thead><tr>{"".join(f"<th>{esc(c['label'])}</th>" for c in columns)}</tr></thead><tbody>{body}</tbody></table>
<div class="hash">شناسه کنترل (SHA-256): <span>{digest}</span></div>
</body></html>"""
	return HTML(string=document).write_pdf()
