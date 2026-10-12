"""Excel import (ورود اکسلی): for every screen where the admin can add something there is a
sample workbook to download and an import that reads the filled workbook back.

The sample has one sheet to fill (a header row, formats already applied to the columns: text
columns stay text so "1405/07/14" is never turned into a Gregorian date, drop-down lists for
choices and for things that already exist), a guide sheet that says what each column wants,
and a hidden sheet with the lists. The import is all or nothing: every row goes through the same
code as the form (same validation, audit and recalculation), and if any row is wrong nothing is
saved and each wrong row is reported with its Excel row number.
"""

from __future__ import annotations

import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from chronicle.calendar import service as calendar_service
from chronicle.db.models import (
	ApprovalProcess,
	AttendanceDevice,
	Department,
	Designation,
	Employee,
	LeaveType,
	Province,
	Shift,
	WorkLocation,
	WorkPattern,
)
from chronicle.engine.rules import PARAMETERS, baseline_rule_set
from chronicle.i18n import _, translate
from chronicle.jalali.dates import (
	WEEKDAY_NAMES,
)
from chronicle.services import cleaning, devices, employees, resources
from chronicle.services.errors import ChronicleError, NotFound, ValidationError

MAX_ROWS = 5000
FORMATTED_ROWS = 1000  # rows of the sample sheet that carry formats and drop-downs

TEXT, NUMBER, BOOL, JDATE, TIME, CHOICE, REF = "text", "number", "bool", "jdate", "time", "choice", "ref"


@dataclass(frozen=True)
class Col:
	key: str
	label: str  # English source text, translated into the header
	kind: str = TEXT
	required: bool = False
	choices: tuple[tuple[str, str], ...] = ()  # (code, English label)
	ref: str | None = None  # a lookup name, see LOOKUPS
	example: str = ""
	help: str = ""


@dataclass
class Spec:
	name: str
	title: str  # English source text
	columns: tuple[Col, ...]
	save: Callable[[resources.Context, dict[str, Any]], Any]
	prepare: Callable[[resources.Context, dict[str, Any], Session], dict[str, Any]] | None = None
	note: str = ""  # shown in the guide sheet (English source text)


# what an "existing thing" column is matched on, and the model it comes from
LOOKUPS: dict[str, tuple[type, str]] = {
	"employees": (Employee, "code"),
	"work-locations": (WorkLocation, "name"),
	"provinces": (Province, "name"),
	"departments": (Department, "name"),
	"designations": (Designation, "name"),
	"shifts": (Shift, "name"),
	"work-patterns": (WorkPattern, "name"),
	"approval-processes": (ApprovalProcess, "name"),
	"leave-types": (LeaveType, "name"),
	"devices": (AttendanceDevice, "name"),
}


def _options(session: Session, ref: str) -> dict[str, int]:
	model, attribute = LOOKUPS[ref]
	column = getattr(model, attribute)
	return {str(v).strip(): i for i, v in session.execute(select(model.id, column)) if v is not None}


# the specs -----------------------------------------------------------------------------------------

REQUEST_TYPES = (
	("daily_leave", _("Daily leave")),
	("hourly_leave", _("Hourly leave")),
	("mission", _("Mission")),
	("overtime", _("Overtime")),
	("attendance_correction", _("Manual time entry")),
)
TARGETS = (
	("employee", _("One employee")),
	("designation", _("A designation")),
	("department", _("A department")),
	("company", _("The whole company")),
)


def _generic(name: str) -> Callable[[resources.Context, dict[str, Any]], Any]:
	def save(ctx: resources.Context, values: dict[str, Any]) -> Any:
		return resources.save(ctx, resources.resource(name), values)

	return save


def _days(ctx: resources.Context, values: dict[str, Any], session: Session) -> dict[str, Any]:
	"""A work pattern row: seven weekday columns, or a rotation like ``Morning*14, Night*14``."""
	shifts = _options(session, "shifts")

	def shift_id(name: str) -> int:
		if name not in shifts:
			raise ValidationError(_("The shift {0} does not exist"), name)
		return shifts[name]

	days: list[dict[str, Any]] = []
	if values.get("pattern_type") == "rotating":
		cycle = 1
		for part in re.split(r"[,،\n]+", str(values.pop("rotation", "") or "")):
			part = part.strip()
			if not part:
				continue
			name, _star, count = part.rpartition("*")
			if not _star:
				name, count = part, "1"
			if not count.strip().isdigit() or int(count) < 1:
				raise ValidationError(_("The rotation {0} must look like Morning*14"), part)
			for _n in range(int(count)):
				days.append({"cycle_day": cycle, "shift_id": shift_id(name.strip())})
				cycle += 1
		if not days:
			raise ValidationError(_("Enter the rotation, for example Morning*14, Night*14"))
		for i in range(7):
			values.pop(f"weekday_{i}", None)
	else:
		values.pop("rotation", None)
		for i in range(7):
			name = values.pop(f"weekday_{i}", None)
			days.append({"weekday": i, "shift_id": shift_id(name) if name else None})
	return {**values, "days": days}


def _segment_minutes(value: str, start: int, end: int) -> tuple[int, int]:
	"""A clock time -> minutes from the shift day's 00:00: today or tomorrow, whichever is nearer
	to the shift (a 01:00 break of a night shift is tomorrow, a 16:00 window before it today)."""
	h, m = (int(x) for x in value.split(":"))
	t = h * 60 + m

	def distance(x: int) -> int:
		return 0 if start <= x <= end else min(abs(x - start), abs(x - end))

	return min((t, t + 1440), key=distance), t


def _segments(ctx: resources.Context, values: dict[str, Any], session: Session) -> dict[str, Any]:
	"""A shift row: breaks and overtime windows written as ``12:30-13:10, 15:00-15:15``."""
	h1, m1 = (int(x) for x in str(values.get("start_time", "00:00")).split(":")[:2])
	h2, m2 = (int(x) for x in str(values.get("end_time", "00:00")).split(":")[:2])
	start, end = h1 * 60 + m1, h2 * 60 + m2
	end += 1440 if end <= start else 0
	segments = []
	for field_key, kind in (("breaks", "break"), ("overtime_windows", "overtime")):
		text = values.pop(field_key, None)
		if not text:
			continue
		try:
			ranges = cleaning.time_ranges(text)
		except ValueError:
			raise ValidationError(_("Write times like 12:30-13:10, 15:00-15:15 ({0})"), text) from None
		for a, b in ranges:
			first, _t = _segment_minutes(a, start, end)
			last = first + (int(b[:2]) * 60 + int(b[3:]) - first) % 1440
			if last == first:
				last += 1440
			segments.append({"kind": kind, "start_minute": first, "end_minute": last})
	return {**values, "segments": segments}


def _steps(ctx: resources.Context, values: dict[str, Any], session: Session) -> dict[str, Any]:
	"""An approval process row: up to four steps, each ``parent``, ``ceo`` or an employee code."""
	people = _options(session, "employees")
	steps = []
	for i in range(1, 5):
		cell = values.pop(f"step_{i}", None)
		if not cell:
			continue
		if cell in ("parent", "ceo"):
			steps.append({"position": len(steps) + 1, "approver_type": cell})
		elif cell in people:
			steps.append(
				{"position": len(steps) + 1, "approver_type": "specific_person", "approver_id": people[cell]}
			)
		else:
			raise ValidationError(_("Step {0}: {1} is not parent, CEO or an employee code"), i, cell)
	return {**values, "steps": steps}


def _parameters(ctx: resources.Context, values: dict[str, Any], session: Session) -> dict[str, Any]:
	law = baseline_rule_set()
	parameters = {p.name: values.pop(p.name, None) for p in PARAMETERS}
	return {
		**values,
		"parameters": {k: (v if v is not None else getattr(law, k)) for k, v in parameters.items()},
	}


def _report_columns(ctx: resources.Context, values: dict[str, Any], session: Session) -> dict[str, Any]:
	keys = [k.strip() for k in re.split(r"[,،\n]+", str(values.pop("column_keys", "") or "")) if k.strip()]
	return {**values, "columns": [{"field": k} for k in keys]}


def _holiday(ctx: resources.Context, values: dict[str, Any]) -> Any:
	return calendar_service.create_holiday(ctx.session, values, ctx.actor.user_id)


def _device(ctx: resources.Context, values: dict[str, Any]) -> Any:
	return devices.save_device(ctx.session, values, ctx.actor.user_id)


def _device_user(ctx: resources.Context, values: dict[str, Any]) -> Any:
	return devices.save_mapping(ctx.session, values, ctx.actor.user_id)


def _employee(ctx: resources.Context, values: dict[str, Any]) -> Any:
	return employees.save(ctx, values)


def _c(key, label, kind=TEXT, required=False, choices=(), ref=None, example="", help=""):
	return Col(key, label, kind, required, tuple(choices), ref, example, help)


GENDERS = (("male", _("Male")), ("female", _("Female")))
STATUSES = (("active", _("Active")), ("suspended", _("Suspended")), ("left", _("Left")))
CONTRACTS = (("full_time", _("Full time")), ("part_time", _("Part time")), ("hourly", _("Hourly")))
HOLIDAY_TYPES = (
	("official_solar", _("Official (solar)")),
	("official_lunar", _("Official (lunar)")),
	("official_occasional", _("Official (announced)")),
	("office_closure", _("Office closure")),
	("company_internal", _("Company holiday")),
)
SCOPES = (
	("country", _("Whole country")),
	("province", _("One province")),
	("work_location", _("One work location")),
	("company", _("Whole company")),
)
CLOSURES = (
	("paid_holiday", _("Paid holiday")),
	("employer_paid", _("Paid by the employer, counted as work")),
	("workday", _("Normal workday (make up later)")),
)
HOLIDAY_STATUSES = (("confirmed", _("Confirmed")), ("predicted", _("Predicted")))
LEAVE_CATEGORIES = (
	("earned", _("Earned")),
	("sick", _("Sick")),
	("unpaid", _("Unpaid")),
	("marriage_bereavement", _("Marriage or bereavement")),
	("hajj", _("Hajj")),
	("maternity", _("Maternity")),
	("paternity", _("Paternity")),
	("other", _("Other")),
)
DIRECTION_MODES = (
	("from_device", _("From Device")),
	("always_in", _("Always In")),
	("always_out", _("Always Out")),
	("pair_by_order", _("Pair by order")),
)
SOURCES = (("summary", _("One row per person")), ("daily", _("One row per day")))
TIME_UNITS = (
	("hours_minutes", _("Hours:minutes")),
	("decimal_hours", _("Decimal hours")),
	("days", _("Days")),
)
EXPORTS = (("excel", _("Excel")), ("pdf", _("PDF")), ("csv", _("Text (CSV)")))

SPECS: dict[str, Spec] = {
	s.name: s
	for s in (
		Spec(
			"work-locations",
			_("Work locations"),
			(
				_c("name", _("Name"), required=True, example="دفتر مرکزی"),
				_c("province_id", _("Province"), REF, True, ref="provinces", example="تهران"),
				_c("city", _("City"), example="تهران"),
				_c("address", _("Address")),
				_c("is_active", _("Active"), BOOL, example="بله"),
			),
			_generic("work-locations"),
		),
		Spec(
			"departments",
			_("Departments"),
			(_c("name", _("Name"), required=True, example="فروش"),),
			_generic("departments"),
		),
		Spec(
			"designations",
			_("Designations"),
			(
				_c("name", _("Name"), required=True, example="مدیر فروش"),
				_c(
					"parent_id",
					_("Parent designation"),
					REF,
					ref="designations",
					help=_("Put parents above their children in the sheet"),
				),
				_c("hard_and_harmful", _("Hard and harmful work"), BOOL),
				_c("committee_ruling_no", _("Committee ruling number")),
				_c("committee_ruling_date", _("Committee ruling date"), JDATE),
				_c("description", _("Description")),
			),
			_generic("designations"),
			note=_("A designation's parent must already exist or be on an earlier row."),
		),
		Spec(
			"designation-assignments",
			_("Designations held"),
			(
				_c(
					"employee_id",
					_("Employee"),
					REF,
					True,
					ref="employees",
					example="1001",
					help=_("The personnel code"),
				),
				_c("designation_id", _("Designation"), REF, True, ref="designations", example="کارشناس فروش"),
				_c("valid_from", _("From"), JDATE, True, example="1405/01/01"),
				_c("valid_to", _("Until"), JDATE),
				_c("note", _("Note")),
			),
			_generic("designation-assignments"),
		),
		Spec(
			"employees",
			_("Employees"),
			(
				_c("code", _("Personnel code"), required=True, example="1001"),
				_c(
					"detail_code",
					_("Detail code"),
					help=_("The accounting detail code (کد تفصیلی)"),
					example="10020030",
				),
				_c("first_name", _("First name"), required=True, example="سارا"),
				_c("last_name", _("Last name"), required=True, example="کریمی"),
				_c("national_id", _("National code"), example="0012345678"),
				_c("gender", _("Gender"), CHOICE, choices=GENDERS),
				_c("mobile", _("Mobile"), example="09121234567"),
				_c("email", _("Email")),
				_c("start_date", _("Start date"), JDATE, True, example="1405/01/01"),
				_c("end_date", _("End date"), JDATE),
				_c("status", _("Status"), CHOICE, choices=STATUSES, example=_("Active")),
				_c("contract_type", _("Contract"), CHOICE, choices=CONTRACTS),
				_c("work_location_id", _("Work location"), REF, ref="work-locations"),
				_c("department_id", _("Department"), REF, ref="departments"),
				_c(
					"designation_id",
					_("Designation"),
					REF,
					ref="designations",
					help=_("The first designation they hold"),
				),
				_c("designation_from", _("Designation from"), JDATE, help=_("Empty: the start date")),
				_c("birth_date", _("Birth date"), JDATE, help=_("Only for young-worker rules")),
				_c("overtime_consent_on_file", _("Written overtime consent on file"), BOOL),
				_c("overtime_consent_date", _("Overtime consent date"), JDATE),
				_c("biometric_consent_date", _("Biometric consent date"), JDATE),
				_c("uses_card_instead_of_biometrics", _("Uses a card instead of biometrics"), BOOL),
				_c("nursing_until", _("Nursing breaks until"), JDATE),
				_c("username", _("Username"), help=_("Empty: the personnel code")),
				_c(
					"password",
					_("Password"),
					required=True,
					help=_("At least 8 characters; the person can change it"),
				),
				_c("is_admin", _("System admin"), BOOL),
			),
			_employee,
			note=_("Every employee is also a user: the username and password create the login."),
		),
		Spec(
			"shifts",
			_("Shifts"),
			(
				_c("name", _("Name"), required=True, example="اداری"),
				_c(
					"shift_type",
					_("Type"),
					CHOICE,
					True,
					choices=(("fixed", _("Fixed")), ("flexible", _("Flexible"))),
					example=_("Fixed"),
				),
				_c("start_time", _("Start"), TIME, True, example="08:00"),
				_c(
					"end_time",
					_("End"),
					TIME,
					True,
					example="16:30",
					help=_("Earlier than the start: the shift ends the next day"),
				),
				_c("latest_start_time", _("Latest start"), TIME),
				_c("required_minutes", _("Required minutes"), NUMBER),
				_c(
					"breaks",
					_("Breaks"),
					example="12:30-13:10",
					help=_("Unpaid breaks, e.g. 12:30-13:10, 15:00-15:15"),
				),
				_c(
					"overtime_windows",
					_("Overtime windows"),
					example="17:00-19:00",
					help=_("When extra work may become overtime, e.g. 17:00-19:00"),
				),
				_c("break_minutes", _("Break minutes without punching"), NUMBER),
				_c("grace_late_minutes", _("Grace for late arrival (minutes)"), NUMBER),
				_c("grace_early_minutes", _("Grace for early leave (minutes)"), NUMBER),
				_c("overtime_after_minutes", _("Extra work counts after (minutes)"), NUMBER),
				_c("color", _("Colour")),
			),
			_generic("shifts"),
			prepare=_segments,
		),
		Spec(
			"work-patterns",
			_("Work patterns"),
			(
				_c("name", _("Name"), required=True, example="هفتگی اداری"),
				_c(
					"pattern_type",
					_("Type"),
					CHOICE,
					True,
					choices=(("weekly", _("Weekly")), ("rotating", _("Rotating"))),
					example=_("Weekly"),
				),
				_c(
					"weekly_rest_day",
					_("Weekly rest day"),
					NUMBER,
					help=_("0 = Saturday ... 6 = Friday; empty: Friday"),
				),
				_c("cycle_start_date", _("First day of the cycle"), JDATE, help=_("Rotating patterns only")),
				*[
					_c(
						f"weekday_{i}",
						WEEKDAY_NAMES[i],
						REF,
						ref="shifts",
						help=_("Weekly patterns: the shift, empty = day off"),
					)
					for i in range(7)
				],
				_c(
					"rotation",
					_("Rotation"),
					help=_("Rotating patterns: shift*days, e.g. Morning*14, Night*14"),
				),
				_c("description", _("Description")),
				_c("change_reason", _("Reason for the change")),
			),
			_generic("work-patterns"),
			prepare=_days,
			note=_(
				"A weekly pattern fills the seven weekday columns; a rotating pattern fills the rotation column."
			),
		),
		Spec(
			"work-pattern-assignments",
			_("Pattern assignments"),
			(
				_c("target_type", _("Applies to"), CHOICE, True, choices=TARGETS, example=_("One employee")),
				_c(
					"employee_id",
					_("Employee"),
					REF,
					ref="employees",
					help=_("When it applies to one employee"),
				),
				_c(
					"designation_id",
					_("Designation"),
					REF,
					ref="designations",
					help=_("When it applies to a designation"),
				),
				_c(
					"department_id",
					_("Department"),
					REF,
					ref="departments",
					help=_("When it applies to a department"),
				),
				_c("work_pattern_id", _("Work pattern"), REF, True, ref="work-patterns"),
				_c(
					"cycle_offset_days",
					_("Start offset in the cycle (days)"),
					NUMBER,
					help=_("Rotating patterns: shifts this group along the cycle"),
				),
				_c("valid_from", _("From"), JDATE, True, example="1405/01/01"),
				_c("valid_to", _("Until"), JDATE),
				_c("note", _("Note")),
			),
			_generic("work-pattern-assignments"),
		),
		Spec(
			"shift-overrides",
			_("One-day shift changes"),
			(
				_c("employee_id", _("Employee"), REF, True, ref="employees"),
				_c("jdate", _("Date"), JDATE, True, example="1405/07/20"),
				_c("shift_id", _("Shift (empty = day off)"), REF, ref="shifts"),
				_c("reason", _("Reason"), required=True),
			),
			_generic("shift-overrides"),
		),
		Spec(
			"leave-types",
			_("Leave types"),
			(
				_c("name", _("Name"), required=True),
				_c("category", _("Category"), CHOICE, True, choices=LEAVE_CATEGORIES, example=_("Other")),
				_c("is_paid", _("Paid"), BOOL),
				_c("deducts_from_balance", _("Deducted from the leave balance"), BOOL),
				_c("allow_hourly", _("Can be hourly"), BOOL),
				_c("max_days_per_request", _("Most days in one request"), NUMBER),
				_c("once_in_service", _("Once in the whole service"), BOOL),
				_c("rule_limit", _("Limit from the rule set (parameter)")),
				_c("description", _("Description")),
			),
			_generic("leave-types"),
		),
		Spec(
			"approval-processes",
			_("Approval processes"),
			(
				_c("name", _("Name"), required=True),
				_c("description", _("Description")),
				_c("step_1", _("Step 1"), help=_("parent, ceo, or the personnel code of a specific person")),
				_c("step_2", _("Step 2"), help=_("parent, ceo, or the personnel code of a specific person")),
				_c("step_3", _("Step 3"), help=_("parent, ceo, or the personnel code of a specific person")),
				_c("step_4", _("Step 4"), help=_("parent, ceo, or the personnel code of a specific person")),
			),
			_generic("approval-processes"),
			prepare=_steps,
		),
		Spec(
			"approval-overrides",
			_("Approval exceptions"),
			(
				_c("employee_id", _("Employee"), REF, True, ref="employees"),
				_c(
					"request_type",
					_("Request type"),
					CHOICE,
					True,
					choices=REQUEST_TYPES,
					example=_("Overtime"),
				),
				_c("process_id", _("Approval process"), REF, True, ref="approval-processes"),
			),
			_generic("approval-overrides"),
		),
		Spec(
			"rule-sets",
			_("Labor rule sets"),
			(
				_c("name", _("Name"), required=True),
				_c("effective_from", _("Effective from"), JDATE, True, example="1406/01/01"),
				_c("description", _("Description")),
				_c("change_reason", _("Reason for the change")),
				*[
					_c(p.name, p.name, NUMBER, help=f"Article {p.article}; empty = the legal value")
					for p in PARAMETERS
				],
			),
			_generic("rule-sets"),
			prepare=_parameters,
			note=_("A value may be better for the worker than the law, never worse."),
		),
		Spec(
			"report-templates",
			_("Report templates"),
			(
				_c("name", _("Name"), required=True),
				_c("source", _("Rows"), CHOICE, True, choices=SOURCES, example=_("One row per person")),
				_c("time_unit", _("Show times as"), CHOICE, choices=TIME_UNITS),
				_c("export_format", _("Format"), CHOICE, choices=EXPORTS),
				_c(
					"column_keys",
					_("Columns"),
					required=True,
					help=_("Field keys separated by commas, e.g. employee, employee_name, worked_minutes"),
				),
				_c("description", _("Description")),
			),
			_generic("report-templates"),
			prepare=_report_columns,
		),
		Spec(
			"holidays",
			_("Holidays"),
			(
				_c("jdate", _("Date"), JDATE, True, example="1405/07/20"),
				_c("title", _("Title"), required=True, example="تعطیلی شرکت"),
				_c(
					"holiday_type",
					_("Type"),
					CHOICE,
					True,
					choices=HOLIDAY_TYPES,
					example=_("Company holiday"),
				),
				_c("scope", _("Applies to"), CHOICE, choices=SCOPES, example=_("Whole company")),
				_c("province_id", _("Province"), REF, ref="provinces"),
				_c("work_location_id", _("Work location"), REF, ref="work-locations"),
				_c("closure_pay_mode", _("How is it paid?"), CHOICE, choices=CLOSURES),
				_c("status", _("Status"), CHOICE, choices=HOLIDAY_STATUSES),
				_c("note", _("Note")),
			),
			_holiday,
		),
		Spec(
			"devices",
			_("Attendance devices"),
			(
				_c("name", _("Name"), required=True, example="VF600 ورودی"),
				_c("adapter_type", _("Device Type"), required=True, example="zkteco"),
				_c("work_location_id", _("Work location"), REF, ref="work-locations"),
				_c("host", _("IP Address"), example="192.168.1.201"),
				_c("port", _("Port"), NUMBER, example="4370"),
				_c("comm_key", _("Comm key")),
				_c("direction_mode", _("Punch Direction"), CHOICE, choices=DIRECTION_MODES),
				_c("polling_interval_minutes", _("Read Every (Minutes)"), NUMBER),
				_c("timeout_seconds", _("Timeout (Seconds)"), NUMBER),
				_c("use_udp", _("Use UDP"), BOOL),
				_c("enabled", _("Enabled"), BOOL),
			),
			_device,
		),
		Spec(
			"device-users",
			_("Device users"),
			(
				_c("device_id", _("Device"), REF, True, ref="devices"),
				_c("device_user_id", _("Device user id"), required=True, example="12"),
				_c("device_user_name", _("Name on the device")),
				_c("employee_id", _("Employee"), REF, ref="employees"),
				_c("valid_from", _("From"), JDATE),
				_c("valid_to", _("Until"), JDATE),
			),
			_device_user,
		),
	)
}


def spec(name: str) -> Spec:
	try:
		return SPECS[name]
	except KeyError:
		raise NotFound(_("Not found")) from None


def available() -> list[dict[str, str]]:
	return [{"name": s.name, "title": s.title} for s in SPECS.values()]


# the sample workbook --------------------------------------------------------------------------------


def _choice_label(choice: tuple[str, str], language: str) -> str:
	return translate(choice[1], language)


def _format_of(col: Col, language: str) -> str:
	return {
		TEXT: translate(_("Text"), language),
		NUMBER: translate(_("Number"), language),
		BOOL: translate(_("Yes or no"), language),
		JDATE: translate(_("Jalali date, e.g. 1405/07/14"), language),
		TIME: translate(_("Time, e.g. 08:30"), language),
		CHOICE: translate(_("Choose from the list"), language),
		REF: translate(_("Choose from the list"), language),
	}[col.kind]


def template(session: Session, name: str, language: str = "fa") -> bytes:
	"""The sample workbook for one import."""
	s = spec(name)
	t = lambda text: translate(text, language)  # noqa: E731
	wb = Workbook()
	data = wb.active
	data.title = t(_("Data entry"))[:31]
	guide = wb.create_sheet(t(_("Guide"))[:31])
	lists = wb.create_sheet("lists")
	lists.sheet_state = "hidden"
	for sheet in (data, guide):
		sheet.sheet_view.rightToLeft = True
	header_fill = PatternFill("solid", fgColor="0F766E")
	required_fill = PatternFill("solid", fgColor="B45309")

	list_column = 0
	for index, col in enumerate(s.columns, start=1):
		letter = get_column_letter(index)
		cell = data.cell(row=1, column=index, value=t(col.label) + (" *" if col.required else ""))
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = required_fill if col.required else header_fill
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
		data.column_dimensions[letter].width = max(14, min(34, len(str(cell.value)) + 6))
		span = f"{letter}2:{letter}{FORMATTED_ROWS + 1}"
		options: list[str] = []
		if col.kind == CHOICE:
			options = [_choice_label(c, language) for c in col.choices]
		elif col.kind == REF and col.ref:
			options = sorted(_options(session, col.ref))
		elif col.kind == BOOL:
			options = [t(_("Yes")), t(_("No"))]
		if options:
			list_column += 1
			for i, option in enumerate(options, start=1):
				lists.cell(row=i, column=list_column, value=option)
			ref = (
				f"lists!${get_column_letter(list_column)}$1:${get_column_letter(list_column)}${len(options)}"
			)
			validation = DataValidation(
				type="list", formula1=ref, allow_blank=not col.required, showErrorMessage=True
			)
			validation.error = t(_("Choose a value from the list"))
			data.add_data_validation(validation)
			validation.add(span)
		elif col.kind == NUMBER:
			validation = DataValidation(
				type="decimal",
				operator="greaterThanOrEqual",
				formula1="0",
				allow_blank=True,
				showErrorMessage=True,
			)
			validation.error = t(_("Enter a number"))
			data.add_data_validation(validation)
			validation.add(span)
		elif col.kind == JDATE:
			validation = DataValidation(
				type="custom",
				formula1=f'=OR({letter}2="",AND(LEN({letter}2)=10,MID({letter}2,5,1)="/",MID({letter}2,8,1)="/"))',
				allow_blank=True,
				showErrorMessage=True,
			)
			validation.error = t(_("Enter the date as 1405/07/14"))
			data.add_data_validation(validation)
			validation.add(span)
		elif col.kind == TIME:
			validation = DataValidation(
				type="custom",
				formula1=f'=OR({letter}2="",AND(LEN({letter}2)=5,MID({letter}2,3,1)=":"))',
				allow_blank=True,
				showErrorMessage=True,
			)
			validation.error = t(_("Enter the time as 08:30"))
			data.add_data_validation(validation)
			validation.add(span)
		# text-like columns stay text, so nothing is turned into a date or a number by Excel
		fmt = "General" if col.kind == NUMBER else "@"
		for row in range(2, FORMATTED_ROWS + 2):
			data.cell(row=row, column=index).number_format = fmt
	data.row_dimensions[1].height = 34
	data.freeze_panes = "A2"

	guide["A1"] = t(s.title)
	guide["A1"].font = Font(bold=True, size=14)
	guide["A2"] = t(
		"Fill the first sheet from row 2. Columns marked * are required. Do not change the header row."
	)
	guide["A3"] = t(s.note) if s.note else ""
	headers = [
		t(_("Column")),
		t(_("Required")),
		t(_("Format")),
		t(_("Allowed values")),
		t(_("Example")),
		t(_("Note")),
	]
	for i, h in enumerate(headers, start=1):
		cell = guide.cell(row=5, column=i, value=h)
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = header_fill
	for r, col in enumerate(s.columns, start=6):
		allowed = ""
		if col.kind == CHOICE:
			allowed = "، ".join(_choice_label(c, language) for c in col.choices)
		elif col.kind == BOOL:
			allowed = f"{t('Yes')}، {t('No')}"
		values = [
			t(col.label),
			t(_("Yes")) if col.required else "",
			_format_of(col, language),
			allowed,
			t(col.example) if col.example else "",
			t(col.help) if col.help else "",
		]
		for i, v in enumerate(values, start=1):
			guide.cell(row=r, column=i, value=v).alignment = Alignment(wrap_text=True, vertical="top")
	for letter, width in zip("ABCDEF", (28, 10, 28, 40, 18, 46), strict=True):
		guide.column_dimensions[letter].width = width
	out = io.BytesIO()
	wb.save(out)
	return out.getvalue()


# reading it back --------------------------------------------------------------------------------------
#
# A workbook does not have to be the sample: the admin uploads any sheet (or a CSV file), the
# panel shows its columns with suggested matches, the admin connects each field to a column
# (or gives one value for every row), and a dry run shows every row as it will be saved:
# times standardised, dates read, Arabic letters replaced, and each problem by row. Only then
# is it imported, all or nothing.

PREVIEW_ROWS = 200


@dataclass
class RowError:
	row: int
	message: str


@dataclass
class RowPreview:
	row: int
	values: dict[str, Any]
	notes: list[str]
	errors: list[str]


@dataclass
class Result:
	created: int = 0
	errors: list[RowError] = field(default_factory=list)
	rows: list[RowPreview] = field(default_factory=list)
	checked: int = 0
	dry_run: bool = False

	@property
	def imported(self) -> bool:
		return not self.dry_run and not self.errors and self.created > 0

	def as_dict(self) -> dict[str, Any]:
		return {
			"imported": self.imported,
			"dry_run": self.dry_run,
			"created": self.created,
			"checked": self.checked,
			"errors": [{"row": e.row, "message": e.message} for e in self.errors],
			"rows": [
				{"row": r.row, "values": r.values, "notes": r.notes, "errors": r.errors} for r in self.rows
			],
		}


@dataclass
class Table:
	"""The cells of the chosen sheet: a header row and the rows under it."""

	sheets: list[str]
	sheet: str
	header_row: int
	headers: list[str]
	rows: list[tuple[int, list[Any]]]  # (row number in the file, cells)


def _read_csv(content: bytes) -> list[list[Any]]:
	import csv

	for encoding in ("utf-8-sig", "cp1256", "latin-1"):
		try:
			decoded = content.decode(encoding)
			break
		except UnicodeDecodeError:
			continue
	sample = decoded[:4096]
	try:
		dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
	except csv.Error:
		dialect = csv.excel
	return [list(r) for r in csv.reader(io.StringIO(decoded), dialect)]


def read_table(
	content: bytes, filename: str = "", sheet: str | None = None, header_row: int | None = None
) -> Table:
	"""The cells of a workbook (.xlsx) or a CSV file. The header row is the first row with text in
	it unless the admin picks another."""
	if filename.lower().endswith((".csv", ".txt")):
		sheets, chosen, grid = ["CSV"], "CSV", _read_csv(content)
	else:
		try:
			workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
		except Exception:
			raise ValidationError(_("The file is not an Excel workbook (.xlsx) or a CSV file")) from None
		sheets = [ws.title for ws in workbook.worksheets if ws.sheet_state == "visible"] or [
			workbook.worksheets[0].title
		]
		chosen = sheet if sheet in sheets else sheets[0]
		grid = [list(r) for r in workbook[chosen].iter_rows(values_only=True)]
		workbook.close()
	if header_row is None:
		header_row = next(
			(i for i, r in enumerate(grid, start=1) if sum(1 for c in r if cleaning.cell_text(c)) >= 1), 1
		)
	if header_row < 1 or header_row > max(len(grid), 1):
		raise ValidationError(_("The header row {0} is outside the sheet"), header_row)
	headers = [cleaning.cell_text(c) for c in (grid[header_row - 1] if grid else [])]
	while headers and not headers[-1]:
		headers.pop()
	rows = [
		(number, cells)
		for number, cells in enumerate(grid[header_row:], start=header_row + 1)
		if any(cleaning.cell_text(c) for c in cells)
	]
	return Table(sheets, chosen, header_row, headers, rows)


def _labels(col: Col, language: str) -> list[str]:
	return [translate(col.label, language), col.label, col.key, *ALIASES.get(col.key, ())]


# other names people give these columns in their own sheets
ALIASES: dict[str, tuple[str, ...]] = {
	"code": ("کد پرسنلی", "شماره پرسنلی", "کد کارمند", "personnel no", "emp code", "employee code"),
	"first_name": ("نام", "first name"),
	"last_name": ("نام خانوادگی", "فامیل", "family", "surname"),
	"national_id": ("کد ملی", "شماره ملی", "national code"),
	"detail_code": ("کد تفصیلی", "تفصیلی"),
	"mobile": ("موبایل", "تلفن همراه", "شماره همراه", "mobile no"),
	"employee_id": ("کد پرسنلی", "شماره پرسنلی", "کارمند", "personnel code"),
	"start_time": ("ساعت ورود", "شروع", "ورود", "from"),
	"end_time": ("ساعت خروج", "پایان", "خروج", "to"),
	"jdate": ("تاریخ", "date", "روز"),
	"valid_from": ("از تاریخ", "تاریخ شروع", "from date"),
	"valid_to": ("تا تاریخ", "تاریخ پایان", "to date"),
	"start_date": ("تاریخ استخدام", "تاریخ شروع کار"),
}


def suggest(s: Spec, headers: list[str], language: str) -> dict[str, int | None]:
	"""Each field's best-matching column (1-based), or None. A column is used once."""
	scored = []
	for col in s.columns:
		for index, header in enumerate(headers, start=1):
			if not header:
				continue
			score = max(
				cleaning.similarity(header.removesuffix("*"), label) for label in _labels(col, language)
			)
			if score >= 0.72:
				scored.append((score, -s.columns.index(col), index, col.key))
	result: dict[str, int | None] = {c.key: None for c in s.columns}
	used: set[int] = set()
	for _score, _order, index, field_key in sorted(scored, reverse=True):
		if result[field_key] is None and index not in used:
			result[field_key] = index
			used.add(index)
	return result


def _samples(table: Table, index: int, count: int = 4) -> list[str]:
	out = []
	for _number, cells in table.rows:
		value = cells[index - 1] if index - 1 < len(cells) else None
		if isinstance(value, datetime | date | time):
			out.append(value.isoformat())
		elif cleaning.cell_text(value):
			out.append(cleaning.cell_text(value))
		if len(out) >= count:
			break
	return out


def _column_view(col: Col, language: str) -> dict[str, Any]:
	return {
		"key": col.key,
		"label": translate(col.label, language),
		"kind": col.kind,
		"required": col.required,
		"help": translate(col.help, language) if col.help else "",
		"example": translate(col.example, language) if col.example else "",
		"choices": [translate(label, language) for _code, label in col.choices],
	}


def inspect_file(
	session: Session,
	name: str,
	content: bytes,
	filename: str,
	language: str,
	sheet: str | None = None,
	header_row: int | None = None,
) -> dict[str, Any]:
	"""What the column-mapping panel needs: the file's sheets and columns with sample values,
	the import's fields, and a suggested match for each field."""
	s = spec(name)
	table = read_table(content, filename, sheet, header_row)
	return {
		"sheets": table.sheets,
		"sheet": table.sheet,
		"header_row": table.header_row,
		"row_count": len(table.rows),
		"columns": [
			{"index": i, "header": h, "samples": _samples(table, i)}
			for i, h in enumerate(table.headers, start=1)
		],
		"fields": [_column_view(c, language) for c in s.columns],
		"suggested": suggest(s, table.headers, language),
	}


def _convert(
	col: Col, raw: Any, session: Session, language: str, cache: dict[str, dict[str, int]]
) -> tuple[Any, str | None]:
	"""A cell -> the field's value and a note when it was standardised."""
	label = translate(col.label, language)
	try:
		if col.kind == TEXT:
			return cleaning.text(raw)
		if col.kind == NUMBER:
			return cleaning.number(raw)
		if col.kind == BOOL:
			value, note = cleaning.boolean(raw)
			if value is None and cleaning.plain(raw):
				raise ValueError(raw)
			return value, note
		if col.kind == JDATE:
			return cleaning.jdate(raw)
		if col.kind == TIME:
			return cleaning.clock(raw)
	except ValueError:
		message = {
			NUMBER: _("{0} must be a number"),
			BOOL: _("{0}: enter Yes or No"),
			JDATE: _("{0}: {1} is not a valid Jalali date"),
			TIME: _("{0}: enter the time as HH:MM"),
		}.get(col.kind, _("{0}: the value {1} is not valid"))
		raise ValidationError(message, label, cleaning.cell_text(raw)) from None
	text = cleaning.cell_text(raw)
	if text == "":
		return None, None
	if col.kind == CHOICE:
		wanted = cleaning.key(text)
		for code, choice in col.choices:
			if wanted in (
				cleaning.key(code),
				cleaning.key(choice),
				cleaning.key(translate(choice, language)),
			):
				return code, None
		raise ValidationError(_("{0}: {1} is not one of the allowed values"), label, text)
	if col.kind == REF:
		if col.ref not in cache:
			cache[col.ref] = {cleaning.key(k): v for k, v in _options(session, col.ref).items()}
		found = cache[col.ref].get(cleaning.key(text))
		if found is None:
			raise ValidationError(_("{0}: {1} does not exist"), label, text)
		return found, None
	return text, None


def _display(col: Col, value: Any, raw: Any, language: str) -> Any:
	if value is None:
		return None
	if col.kind == CHOICE:
		return next((translate(label, language) for code, label in col.choices if code == value), value)
	if col.kind == REF:
		return cleaning.cell_text(raw)
	if col.kind == BOOL:
		return translate(_("Yes") if value else _("No"), language)
	return value


def _header_map(table: Table, s: Spec, language: str) -> dict[str, int]:
	"""The sample's own headers: each field to the column with its exact label."""
	by_name: dict[str, Col] = {}
	for col in s.columns:
		for name in (translate(col.label, language), col.label, col.key):
			by_name[cleaning.key(name)] = col
	mapping: dict[str, int] = {}
	for index, header in enumerate(table.headers, start=1):
		col = by_name.get(cleaning.key(header.removesuffix("*")))
		if col is not None and col.key not in mapping:
			mapping[col.key] = index
	return mapping


Mapping = dict[str, Any]  # field key -> column index (1-based) or {"value": "..."} for every row


def _check_mapping(s: Spec, mapping: Mapping, table: Table, language: str) -> dict[str, Any]:
	keys = {c.key for c in s.columns}
	clean: dict[str, Any] = {}
	for field_key, source in (mapping or {}).items():
		if field_key not in keys or source in (None, "", 0):
			continue
		if isinstance(source, dict):
			if cleaning.plain(source.get("value")):
				clean[field_key] = {"value": source.get("value")}
			continue
		try:
			index = int(source)
		except (TypeError, ValueError):
			raise ValidationError(_("The column {0} is not in the file"), source) from None
		if not 1 <= index <= max(len(table.headers), 1):
			raise ValidationError(_("The column {0} is not in the file"), index)
		clean[field_key] = index
	missing = [translate(c.label, language) for c in s.columns if c.required and c.key not in clean]
	if missing:
		raise ValidationError(_("Connect a column to {0}"), "، ".join(missing))
	return clean


def run(
	ctx: resources.Context,
	name: str,
	content: bytes,
	*,
	filename: str = "",
	mapping: Mapping | None = None,
	sheet: str | None = None,
	header_row: int | None = None,
	dry_run: bool = False,
) -> Result:
	"""Import a sheet. Nothing is saved unless every row is valid; a dry run saves nothing but
	goes through the same checks and returns every row as it would be saved."""
	s = spec(name)
	language = ctx.actor.language
	table = read_table(content, filename, sheet, header_row)
	if mapping is None:
		mapping = _header_map(table, s, language)
		missing = [translate(c.label, language) for c in s.columns if c.required and c.key not in mapping]
		if missing:
			raise ValidationError(_("The sheet has no column {0}"), "، ".join(missing))
	sources = _check_mapping(s, mapping, table, language)
	columns = {c.key: c for c in s.columns}
	if len(table.rows) > MAX_ROWS:
		raise ValidationError(_("The sheet has more than {0} rows"), MAX_ROWS)
	result = Result(dry_run=dry_run)
	cache: dict[str, dict[str, int]] = {}
	outer = ctx.session.begin_nested()
	try:
		for row_number, cells in table.rows:
			inner = ctx.session.begin_nested()
			values: dict[str, Any] = {}
			shown: dict[str, Any] = {}
			notes: list[str] = []
			problems: list[str] = []
			for field_key, source in sources.items():
				col = columns[field_key]
				if isinstance(source, dict):
					raw = source["value"]
				else:
					raw = cells[source - 1] if source - 1 < len(cells) else None
				try:
					value, note = _convert(col, raw, ctx.session, language, cache)
				except ChronicleError as e:
					problems.append(e.translated(language))
					continue
				if note:
					notes.append(f"{translate(col.label, language)}: {note}")
				if value is None and col.required:
					problems.append(
						translate(_("{0} is required"), language).format(translate(col.label, language))
					)
				if value is not None:
					values[field_key] = value
					shown[field_key] = _display(col, value, raw, language)
			try:
				if problems:
					raise ValidationError("; ".join(problems).replace("{", "{{").replace("}", "}}"))
				if s.prepare:
					values = s.prepare(ctx, values, ctx.session)
				s.save(ctx, values)
				inner.commit()
				cache.clear()  # what this row created can be used by the rows below it
				result.created += 1
			except (ChronicleError, IntegrityError) as e:
				inner.rollback()
				problems = [
					e.translated(language)
					if isinstance(e, ChronicleError)
					else translate(_("A value is not valid ({0})"), language).format(
						getattr(getattr(e.orig, "diag", None), "constraint_name", "")
					)
				]
				result.errors.append(RowError(row_number, problems[0]))
			result.checked += 1
			if len(result.rows) < PREVIEW_ROWS or problems:
				result.rows.append(RowPreview(row_number, shown, notes, problems))
	except Exception:
		outer.rollback()
		raise
	if result.errors or dry_run:
		outer.rollback()
		if result.errors:
			result.created = 0
	else:
		outer.commit()
	return result
