import io

import pytest
from openpyxl import load_workbook

from hozur.reports import catalog
from hozur.reports.exporters import (
	DAYS_UNIT,
	DECIMAL_HOURS,
	FIXED_WIDTH,
	HOURS_MINUTES,
	content_hash,
	display_value,
	hours_minutes,
	to_text,
	to_xlsx,
)

COLUMNS = [
	{"key": "employee", "label": "کد پرسنلی", "kind": "text", "width": 6},
	{"key": "employee_name", "label": "نام", "kind": "text", "width": 14},
	{"key": "worked_minutes", "label": "کارکرد", "kind": "minutes", "width": 8},
	{"key": "leave_days", "label": "مرخصی", "kind": "days", "width": 6},
	{"key": "absent_days", "label": "غیبت", "kind": "count", "width": 4},
]
ROWS = [
	{
		"employee": "101",
		"employee_name": "آراد فراب",
		"worked_minutes": 10580,
		"leave_days": 1.25,
		"absent_days": 0,
	},
	{
		"employee": "102",
		"employee_name": "سارا یکتا",
		"worked_minutes": 9645,
		"leave_days": 0,
		"absent_days": 1,
	},
]


def test_time_units():
	assert hours_minutes(10580) == "176:20"
	assert display_value(90, "minutes", HOURS_MINUTES, 440) == "1:30"
	assert display_value(90, "minutes", DECIMAL_HOURS, 440) == "1.50"
	assert display_value(880, "minutes", DAYS_UNIT, 440) == "2.00"
	assert display_value(1.25, "days", HOURS_MINUTES, 440) == "1.25"
	assert display_value(None, "minutes", HOURS_MINUTES, 440) == ""


def test_excel_is_rtl_numeric_and_totalled():
	content = to_xlsx("کارکرد", ["مهر ۱۴۰۵"], COLUMNS, ROWS, HOURS_MINUTES, 440)
	ws = load_workbook(io.BytesIO(content)).active
	assert ws.sheet_view.rightToLeft
	header = next(r for r in ws.iter_rows(values_only=True) if r and r[0] == "کد پرسنلی")
	assert list(header) == [c["label"] for c in COLUMNS]
	data = next(r for r in ws.iter_rows() if r[0].value == "101")
	# a real duration Excel can add up (openpyxl reads [h]:mm cells back as timedelta)
	assert data[2].number_format == "[h]:mm"
	assert data[2].value.total_seconds() == 10580 * 60
	total = next(r for r in ws.iter_rows(values_only=True) if r and r[0] == "جمع")
	assert total[2].total_seconds() == (10580 + 9645) * 60 and total[4] == 1


def test_text_delimiters_and_ascii_digits():
	text = to_text(COLUMNS, ROWS, DECIMAL_HOURS, 440, "Tab").decode("utf-8")
	lines = text.splitlines()
	assert lines[1].split("\t") == ["101", "آراد فراب", "176.33", "1.25", "0"]
	pipe = to_text(COLUMNS, ROWS, HOURS_MINUTES, 440, "Pipe", header=False).decode("utf-8")
	assert pipe.splitlines()[0] == "101|آراد فراب|176:20|1.25|0"


def test_fixed_width_text():
	text = to_text(COLUMNS, ROWS, HOURS_MINUTES, 440, FIXED_WIDTH, header=False).decode("utf-8")
	first = text.splitlines()[0]
	assert len(first) == sum(c["width"] for c in COLUMNS)
	assert first.startswith("101   آراد فراب     176:20  1.25  0   ")


def test_windows_1256_uses_arabic_yeh_for_old_payroll_software():
	content = to_text(COLUMNS, ROWS, HOURS_MINUTES, 440, "Comma", "Windows-1256")
	decoded = content.decode("cp1256")
	assert "سارا يکتا" in decoded  # Persian yeh became Arabic yeh, keheh stays
	assert "?" not in decoded


def test_utf8_with_bom():
	assert to_text(COLUMNS, ROWS, HOURS_MINUTES, 440, "Comma", "UTF-8 with BOM").startswith(b"\xef\xbb\xbf")


def test_hash_detects_changes():
	h = content_hash(COLUMNS, ROWS)
	assert h == content_hash(COLUMNS, [dict(r) for r in ROWS])
	changed = [dict(ROWS[0], worked_minutes=10581), ROWS[1]]
	assert h != content_hash(COLUMNS, changed)


def test_pdf_renders(tmp_path):
	pytest.importorskip("weasyprint")
	from pathlib import Path

	from hozur.reports.exporters import to_pdf

	fonts = Path(__file__).resolve().parents[2] / "hozur" / "public" / "fonts" / "vazirmatn"
	pdf = to_pdf("کارکرد ماهانه", ["مهر ۱۴۰۵"], COLUMNS, ROWS, HOURS_MINUTES, 440, f"file://{fonts}")
	assert pdf.startswith(b"%PDF") and len(pdf) > 5000


def test_catalog_accepts_only_known_fields():
	assert catalog.resolve("worked_minutes", catalog.SUMMARY, set()).kind == catalog.MINUTES
	assert catalog.resolve("leave:استحقاقی", catalog.DAILY, {"استحقاقی"}).kind == catalog.DAYS
	with pytest.raises(ValueError):
		catalog.resolve("salary", catalog.SUMMARY, set())
	with pytest.raises(ValueError):
		catalog.resolve("jdate", catalog.SUMMARY, set())  # daily-only field
	with pytest.raises(ValueError):
		catalog.resolve("leave:unknown", catalog.SUMMARY, {"استحقاقی"})
