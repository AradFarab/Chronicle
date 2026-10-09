"""Excel import: a sample workbook for every create form, and an all-or-nothing import."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook
from orgfactory import TODAY
from sqlalchemy import func, select

from hozur.db.models import Designation, Employee, Holiday, Shift, WorkLocation, WorkPattern
from hozur.services import excel, resources
from hozur.services.errors import ValidationError


def ctx(org, actor=None):
	return resources.Context(org.db, actor or org.admin, TODAY)


def workbook(headers, *rows) -> bytes:
	wb = Workbook()
	ws = wb.active
	ws.append(headers)
	for row in rows:
		ws.append(list(row))
	out = io.BytesIO()
	wb.save(out)
	return out.getvalue()


def labels(name):
	return [excel.translate(c.label, "fa") + (" *" if c.required else "") for c in excel.spec(name).columns]


def test_every_import_has_a_sample_with_text_formats_and_dropdowns(org):
	for name in excel.SPECS:
		wb = load_workbook(io.BytesIO(excel.template(org.db, name)))
		data = wb.worksheets[0]
		columns = excel.spec(name).columns
		assert [c.value for c in data[1]] == labels(name), name
		for index, col in enumerate(columns, start=1):
			fmt = data.cell(row=2, column=index).number_format
			assert fmt == ("General" if col.kind == excel.NUMBER else "@"), (name, col.key)
		assert len(wb.worksheets) == 3 and wb["lists"].sheet_state == "hidden"
		validated = {str(rng) for dv in data.data_validations.dataValidation for rng in dv.sqref.ranges}
		for index, col in enumerate(columns, start=1):
			if col.kind not in (excel.TEXT, excel.REF):  # a pick-list of things that exist (maybe none yet)
				letter = data.cell(row=1, column=index).column_letter
				assert any(r.startswith(f"{letter}2:") for r in validated), (name, col.key)


def test_the_sample_covers_every_field_of_its_form(org):
	"""A field added to a form must be added to its Excel columns too."""
	skip = {"provinces"}
	for resource in resources.RESOURCES.values():
		if resource.name in skip or resource.name not in excel.SPECS or not resource.fields:
			continue
		keys = {c.key for c in excel.spec(resource.name).columns}
		missing = [
			f
			for f in resource.fields
			if f not in keys
			and f
			not in {
				"days",
				"steps",
				"parameters",
				"columns",
				"shared_roles",
				"group_by",
				"text_delimiter",
				"text_encoding",
				"default_range",
				"is_ceo",  # the CEO designation is created at install; it is renamed in the form
			}
		]
		assert not missing, (resource.name, missing)
	from hozur.services.employees import FIELDS

	assert set(FIELDS) <= {c.key for c in excel.spec("employees").columns}


def test_importing_work_locations_and_designations_in_order(org):
	province = org.db.scalar(select(excel.Province.name).limit(1))
	content = workbook(
		labels("work-locations"),
		("دفتر شیراز", province, "شیراز", None, "بله"),
		("انبار", province, None, None, "خیر"),
	)
	result = excel.run(ctx(org), "work-locations", content)
	assert result.as_dict() == {"imported": True, "created": 2, "errors": []}
	assert org.db.scalar(select(WorkLocation.is_active).where(WorkLocation.name == "انبار")) is False
	# a child may refer to a parent created earlier in the same sheet
	content = workbook(
		labels("designations"),
		("مدیر مالی", org.ceo_designation.name),
		("کارشناس مالی", "مدیر مالی"),
	)
	assert excel.run(ctx(org), "designations", content).created == 2
	child = org.db.scalar(select(Designation).where(Designation.name == "کارشناس مالی"))
	assert child.parent_id == org.db.scalar(select(Designation.id).where(Designation.name == "مدیر مالی"))


def sheet_rows(name, *dicts):
	"""A workbook for `name` with one row per dict of column key -> cell."""
	columns = excel.spec(name).columns
	return workbook(labels(name), *[[d.get(c.key) for c in columns] for d in dicts])


def test_importing_employees_creates_their_logins_and_designations(org):
	content = sheet_rows(
		"employees",
		{
			"code": "2001",
			"detail_code": "30001",
			"first_name": "علی",
			"last_name": "احمدی",
			"start_date": "۱۴۰۵/۰۱/۱۵",
			"designation_id": org.expert_designation.name,
			"password": "long-pass-1",
		},
		{
			"code": "2002",
			"first_name": "مینا",
			"last_name": "رضایی",
			"start_date": "1405/02/01",
			"password": "long-pass-2",
			"username": "mina",
		},
	)
	result = excel.run(ctx(org), "employees", content)
	assert result.as_dict() == {"imported": True, "created": 2, "errors": []}
	ali = org.db.scalar(select(Employee).where(Employee.code == "2001"))
	assert ali.start_date == "1405/01/15" and ali.detail_code == "30001"  # Persian digits accepted
	from hozur.db.models import AppUser

	assert org.db.get(AppUser, ali.user_id).username == "2001"
	mina = org.db.scalar(select(Employee).where(Employee.code == "2002"))
	assert org.db.get(AppUser, mina.user_id).username == "mina"
	from hozur.services import access

	assert access.designations_on(org.db, ali.id, TODAY) == [org.expert_designation.id]


def test_one_bad_row_saves_nothing_and_every_bad_row_is_reported(org):
	before = org.db.scalar(select(func.count()).select_from(Shift))
	content = workbook(
		labels("shifts"),
		("شیفت خوب", "ثابت", "08:00", "16:30") + (None,) * 9,
		("شیفت بد", "ثابت", "25:00", "16:30") + (None,) * 9,
		("شیفت ناقص", "ثابت", "08:00", None) + (None,) * 9,
	)
	result = excel.run(ctx(org), "shifts", content)
	assert not result.imported and result.created == 0
	assert [e.row for e in result.errors] == [3, 4]
	assert org.db.scalar(select(func.count()).select_from(Shift)) == before  # nothing was kept


def test_rotation_columns_make_a_guards_pattern(org):
	for name, start, end in (("صبح", "06:00", "18:00"), ("شب", "18:00", "06:00")):
		org.db.add(
			Shift(
				name=name,
				start_time=__import__("datetime").time.fromisoformat(start),
				end_time=__import__("datetime").time.fromisoformat(end),
			)
		)
	org.db.flush()
	cols = excel.spec("work-patterns").columns
	row = {c.key: None for c in cols}
	row.update(name="نگهبانی", pattern_type="چرخشی", cycle_start_date="1405/07/01", rotation="صبح*14، شب*14")
	content = workbook(labels("work-patterns"), [row[c.key] for c in cols])
	result = excel.run(ctx(org), "work-patterns", content)
	assert result.imported, result.errors
	pattern = org.db.scalar(select(WorkPattern).where(WorkPattern.name == "نگهبانی"))
	assert len(pattern.days) == 28
	assert [d.cycle_day for d in pattern.days] == list(range(1, 29))
	assert pattern.days[0].shift_id != pattern.days[14].shift_id


def test_holidays_and_weekly_patterns(org):
	cols = excel.spec("holidays").columns
	row = {c.key: None for c in cols}
	row.update(jdate="1405/07/20", title="تعطیلی شرکت", holiday_type="تعطیلی شرکت", scope="کل شرکت")
	assert excel.run(ctx(org), "holidays", workbook(labels("holidays"), [row[c.key] for c in cols])).imported
	assert org.db.scalar(select(Holiday.scope).where(Holiday.title == "تعطیلی شرکت")) == "company"


def test_a_file_that_is_not_excel_or_misses_columns_is_refused(org):
	with pytest.raises(ValidationError):
		excel.run(ctx(org), "departments", b"not an excel file")
	with pytest.raises(ValidationError, match="no column"):
		excel.run(ctx(org), "departments", workbook(["something else"], ("x",)))


def test_real_excel_dates_are_refused_with_a_hint(org):
	import datetime

	content = workbook(labels("holidays"), (datetime.date(2026, 10, 12), "x", "تعطیلی شرکت"))
	result = excel.run(ctx(org), "holidays", content)
	assert not result.imported and [e.row for e in result.errors] == [2]
